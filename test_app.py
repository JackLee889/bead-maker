from io import BytesIO
import csv
import hashlib
import json
from pathlib import Path
import tempfile
import zipfile
from unittest.mock import patch

from PIL import Image, ImageDraw
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from brand_palettes import BRAND_PALETTES, load_brand_palette


def sample_png(solid: str | None = None) -> bytes:
    image = Image.new("RGB", (160, 100), solid or "#F3EFE7")
    if solid is None:
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, 79, 99), fill="#D33242")
        draw.ellipse((45, 15, 135, 92), fill="#2B68B8")
        draw.polygon([(80, 5), (150, 55), (95, 95)], fill="#F4C542")
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def test_builtin_palette() -> None:
    app = AppTest.from_file("app.py", default_timeout=60).run()
    assert not app.exception
    app.get("file_uploader")[0].upload(
        "qa_source.png", sample_png(), "image/png"
    ).run(timeout=60)
    app.number_input[0].set_value(40)
    app.selectbox(key="fit_mode__zh-CN").set_value("完整留白")
    app.selectbox(key="generic_palette__zh-CN").set_value("通用 32 色")
    app.selectbox(key="pdf_layout__zh-CN").set_value("30 × 40")
    app.run(timeout=60)
    app.button[0].click().run(timeout=60)
    assert not app.exception
    assert not app.error
    result = app.session_state["result"]
    assert (result["columns"], result["rows"]) == (40, 25)
    assert result["grid"].startswith(b"\x89PNG")
    assert result["preview"].startswith(b"\x89PNG")
    assert result["pdf"].startswith(b"%PDF")
    assert result["bundle"].startswith(b"PK")
    assert result["csv"].startswith(b"\xef\xbb\xbf")
    assert sum(item["count"] for item in result["usage"]) == 1000
    assert result["xlsx"].startswith(b"PK")
    assert len(app.get("download_button")) == 5
    excel = load_workbook(BytesIO(result["xlsx"]))
    assert (excel["拼豆图"].max_row, excel["拼豆图"].max_column) == (25, 40)
    with zipfile.ZipFile(BytesIO(result["bundle"])) as archive:
        assert archive.read("bead_pattern.xlsx") == result["xlsx"]
        settings = json.loads(archive.read("settings.json"))
        assert settings["source_sha256"] == hashlib.sha256(sample_png()).hexdigest()
        assert settings["source_name"] == "qa_source.png"
        assert settings["png_cell_size"] == 30
        assert settings["pdf_cells_per_page"] == [30, 40]
        assert settings["app_version"] == "1.5.0"
        assert settings["line_strength"] == 0.65
        assert settings["sharpness"] == 0.25

    # Ordinary reruns keep valid downloads; changed parameters invalidate them.
    app.run(timeout=60)
    assert len(app.get("download_button")) == 5
    assert app.select_slider[0].value == "适应窗口"
    current_bundle = app.session_state["result"]["bundle"]
    for zoom in (4, 16, "适应窗口"):
        app.select_slider[0].set_value(zoom).run(timeout=60)
        assert not app.exception
        assert not app.error
        assert app.session_state["result"]["bundle"] == current_bundle
    app.slider[0].set_value(1.2).run(timeout=60)
    assert "result" not in app.session_state
    assert not app.get("download_button")
    app.button[0].click().run(timeout=60)
    assert not app.error
    assert len(app.get("download_button")) == 5
    with patch("core.make_pdf", side_effect=ValueError("测试导出失败")):
        app.button[0].click().run(timeout=60)
    assert not app.exception
    assert any("测试导出失败" in error.value for error in app.error)
    assert "result" not in app.session_state
    assert not app.get("download_button")


def test_custom_palette() -> None:
    app = AppTest.from_file("app.py", default_timeout=60).run()
    app.get("file_uploader")[0].upload(
        "solid.png", sample_png("#CC2233"), "image/png"
    ).run(timeout=60)
    app.radio[0].set_value("导入品牌/自定义色号表").run(timeout=60)
    palette_csv = (
        b"brand,code,name,hex\n"
        b"BrandX,X01,Red,#CC2233\n"
        b"BrandX,X02,White,#FFFFFF\n"
        b"BrandY,Y01,Black,#000000\n"
    )
    app.get("file_uploader")[1].upload(
        "palette.csv", palette_csv, "text/csv"
    ).run(timeout=60)
    app.multiselect[0].set_value(["BrandX"]).run(timeout=60)
    app.button[-1].click().run(timeout=60)
    assert not app.exception
    assert not app.error
    assert {item["brand"] for item in app.session_state["result"]["usage"]} == {
        "BrandX"
    }
    excel = load_workbook(BytesIO(app.session_state["result"]["xlsx"]))
    assert excel["拼豆图"]["A1"].value == "X01"
    with zipfile.ZipFile(BytesIO(app.session_state["result"]["bundle"])) as archive:
        settings = json.loads(archive.read("settings.json"))
        assert settings["palette_sha256"] == hashlib.sha256(palette_csv).hexdigest()
        assert settings["selected_brands"] == ["BrandX"]
    # Same filename and brand, different content must invalidate the old result.
    app.get("file_uploader")[1].upload(
        "palette.csv", palette_csv.replace(b"#CC2233", b"#2266CC"), "text/csv"
    ).run(timeout=60)
    assert "result" not in app.session_state
    assert len(app.get("download_button")) == 1  # Only the palette template.
    app.button[0].click().run(timeout=60)
    assert not app.error
    app.multiselect[0].set_value([]).run(timeout=60)
    app.button[0].click().run(timeout=60)
    assert app.error
    assert "result" not in app.session_state


def test_uniform_auto_palette() -> None:
    app = AppTest.from_file("app.py", default_timeout=60).run()
    app.get("file_uploader")[0].upload(
        "solid.png", sample_png("#336699"), "image/png"
    ).run(timeout=60)
    app.radio[0].set_value("自动从原画取色").run(timeout=60)
    app.button[0].click().run(timeout=60)
    assert not app.exception
    assert not app.error
    result = app.session_state["result"]
    usage = result["usage"]
    with zipfile.ZipFile(BytesIO(result["bundle"])) as archive:
        assert json.loads(archive.read("settings.json"))["auto_count"] == 24
    assert len(usage) == 1
    assert usage[0]["count"] == result["columns"] * result["rows"]
    app.checkbox[1].set_value(False).run(timeout=60)
    app.button[0].click().run(timeout=60)
    assert not app.exception
    assert not app.error
    excel = load_workbook(BytesIO(app.session_state["result"]["xlsx"]))
    assert excel["拼豆图"]["A1"].value is None
    assert excel["拼豆图"]["A1"].fill.patternType == "solid"
    app.get("file_uploader")[0].upload("broken.png", b"invalid image", "image/png").run(timeout=60)
    assert not app.exception
    assert app.error
    assert "result" not in app.session_state
    assert not app.get("download_button")


def test_brand_palette_workflow() -> None:
    app = AppTest.from_file("app.py", default_timeout=60).run()
    app.get("file_uploader")[0].upload("brand_sample.png", sample_png(), "image/png").run()
    app.number_input[0].set_value(8)
    app.radio[0].set_value("内置品牌色卡").run()
    assert not app.exception
    assert app.selectbox(key="bead_brand").value == "MARD"

    for brand, versions in BRAND_PALETTES.items():
        for palette_id in versions:
            app.selectbox(key="bead_brand").set_value(brand).run()
            app.selectbox(key="brand_palette_id__zh-CN").set_value(palette_id).run()
            assert not app.exception
            assert "result" not in app.session_state
            app.button[0].click().run(timeout=60)
            assert not app.exception
            assert not app.error
            result = app.session_state["result"]
            palette = load_brand_palette(palette_id)
            valid_codes = {c.code for c in palette.select()}
            assert {item["brand"] for item in result["usage"]} == {brand}
            assert {item["code"] for item in result["usage"]} <= valid_codes
            assert sum(item["count"] for item in result["usage"]) == 40
            assert len(app.get("download_button")) == 5
            excel = load_workbook(BytesIO(result["xlsx"]))
            assert (excel["拼豆图"].max_row, excel["拼豆图"].max_column) == (5, 8)
            assert all(cell.data_type == "s" and cell.value in valid_codes
                       for row in excel["拼豆图"] for cell in row)
            with zipfile.ZipFile(BytesIO(result["bundle"])) as archive:
                assert archive.read("bead_pattern.xlsx") == result["xlsx"]
                settings = json.loads(archive.read("settings.json"))
                assert settings["brand_palette"]["id"] == palette_id
                assert settings["palette_sha256"] == palette.sha256
                assert settings["selected_brands"] == [brand]
                assert set(settings["palette_groups"]) == set(palette.group_counts)
                rows = list(csv.DictReader(archive.read("palette_used.csv").decode("utf-8-sig").splitlines()))
                assert {row["code"] for row in rows} == valid_codes
                notice = archive.read("palette_source.txt").decode("utf-8")
                assert "MIT License" in notice and palette.sha256 in notice

    # Restrict the last brand to one group and verify that every export uses it.
    group = next(iter(palette.group_counts))
    group_widget = f"brand_groups_{palette_id}__zh-CN"
    app.multiselect(key=group_widget).set_value([group]).run()
    assert "result" not in app.session_state
    assert not app.get("download_button")
    app.button[0].click().run(timeout=60)
    assert not app.exception
    result = app.session_state["result"]
    assert {row["code"] for row in result["usage"]} <= {c.code for c in palette.select([group])}
    with zipfile.ZipFile(BytesIO(result["bundle"])) as archive:
        assert json.loads(archive.read("settings.json"))["palette_groups"] == [group]

    # Empty selection cannot generate or retain downloads.
    app.multiselect(key=group_widget).set_value([]).run()
    assert not app.exception
    assert app.button[0].disabled
    assert "result" not in app.session_state
    assert not app.get("download_button")
    app.multiselect(key=group_widget).set_value([group]).run()
    app.button[0].click().run(timeout=60)
    assert not app.error
    # Missing bundled data must not silently fall back to a generic palette.
    with tempfile.TemporaryDirectory() as directory, patch("brand_palettes.DATA_DIR", Path(directory)):
        app.run()
    assert not app.exception
    assert app.error
    assert app.button[0].disabled
    assert "result" not in app.session_state
    assert not app.get("download_button")


def main() -> None:
    test_builtin_palette()
    test_custom_palette()
    test_uniform_auto_palette()
    test_brand_palette_workflow()
    from test_i18n import test_app_language_switch
    test_app_language_switch()
    print("app end-to-end tests passed")


if __name__ == "__main__":
    main()
