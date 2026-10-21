"""Language fallback, translation integrity and output consistency regressions."""
from io import BytesIO
import ast
import base64
import json
from pathlib import Path
from string import Formatter
import unittest
from unittest.mock import patch
import zipfile

import numpy as np
from openpyxl import load_workbook
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from core import BeadColor, builtin_palette, make_pdf
from i18n import LANGUAGES, LocalizedError, Translator, browser_language, catalog, country_language, error_text, resolve_language, valid_preference
from localized_exports import localized_palette


class I18nTests(unittest.TestCase):
    def test_pdf_cjk_font_cmaps_match_character_collections(self):
        palette = [BeadColor('001', '白色', 255, 255, 255)]
        for language, ordering, encoding in (('zh-CN', 'GB1', 'UniGB-UCS2-H'),
                                             ('zh-TW', 'CNS1', 'UniCNS-UCS2-H'),
                                             ('ja', 'Japan1', 'UniJIS-UCS2-H')):
            pdf = PdfReader(BytesIO(make_pdf(np.zeros((1, 1), dtype=np.int16), palette, 5, language=language)))
            fonts = [font.get_object() for font in pdf.pages[0]['/Resources']['/Font'].values()]
            cjk = [font for font in fonts if font.get('/DescendantFonts')]
            self.assertEqual(len(cjk), 1)
            self.assertEqual(cjk[0]['/Encoding'], '/' + encoding)
            self.assertEqual(cjk[0]['/DescendantFonts'][0]['/CIDSystemInfo']['/Ordering'], ordering)
            self.assertIn('白色', pdf.pages[0].extract_text())

    def test_detection_priority_and_fallback(self):
        for country, language in {'CN': 'zh-CN', 'TW': 'zh-TW', 'HK': 'zh-TW', 'MO': 'zh-TW', 'JP': 'ja', 'US': 'en', 'SG': 'en'}.items():
            self.assertEqual(resolve_language('auto', country, 'en-US'), language)
            self.assertEqual(resolve_language('ja', country, 'zh-CN'), 'ja')
        for country in (None, '', 'ZZ', 'XX', 'garbage', {}, ['JP']):
            self.assertIsNone(country_language(country))
            self.assertEqual(resolve_language('auto', country, 'zh-Hant-HK'), 'zh-TW')
        for tag, expected in {'zh': 'zh-CN', 'zh_SG': 'zh-CN', 'zh-Hant': 'zh-TW', 'zh-TW': 'zh-TW', 'ja-JP': 'ja', 'de-DE': 'en', '': 'zh-CN'}.items():
            self.assertEqual(browser_language(tag), expected)
        self.assertIsNone(valid_preference('<script>'))
        self.assertIsNone(valid_preference({}))

    def test_catalogs_cover_messages_and_preserve_placeholders(self):
        base = catalog('zh-CN')
        fields = lambda value: {name for _, name, _, _ in Formatter().parse(value) if name is not None}
        for language in LANGUAGES:
            translated = catalog(language)
            self.assertEqual(base.keys(), translated.keys())
            for key, value in translated.items():
                with self.subTest(language=language, message=key):
                    self.assertTrue(value.strip())
                    self.assertEqual(fields(key), fields(value))
        for filename in ('app.py', 'core.py', 'brand_palettes.py', 'excel_export.py', 'language_ui.py', 'localized_exports.py', 'grid_viewer.py'):
            for node in ast.walk(ast.parse(Path(filename).read_text(encoding='utf-8'))):
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in ('t', 'LocalizedError') and node.args:
                    message = node.args[0]
                    if isinstance(message, ast.Constant) and isinstance(message.value, str):
                        self.assertIn(message.value, base)
        for color in builtin_palette():
            self.assertIn(color.name, base)

    def test_error_arguments_and_imported_labels_remain_literal(self):
        error = LocalizedError('CSV 第 {row_no} 行颜色无效：{exc}', row_no=7,
                               exc=LocalizedError('RGB 无效：{code}', code='白色001'))
        self.assertIn('row 7', error_text(error, 'en'))
        self.assertIn('白色001', error_text(error, 'en'))
        palette = [BeadColor('001', '白色', 255, 255, 255, '我的品牌')]
        for language in LANGUAGES:
            self.assertEqual(localized_palette(palette, '导入品牌/自定义色号表', language), palette)
        original = builtin_palette()
        translated = localized_palette(original, '内置通用色库', 'en')
        self.assertEqual([c.code for c in original], [c.code for c in translated])
        self.assertEqual([c.rgb for c in original], [c.rgb for c in translated])
        self.assertEqual(original[0].name, '纯白')
        self.assertEqual(translated[0].name, 'Pure white')


def test_app_language_switch():
    from test_app import sample_png
    app = AppTest.from_file(str(Path('app.py').resolve()), default_timeout=60).run()
    app.get('file_uploader')[0].upload('language.png', sample_png(), 'image/png').run()
    app.number_input(key='columns').set_value(13)
    app.checkbox(key='keep_ratio').set_value(False).run()
    app.number_input(key='rows').set_value(9)
    app.selectbox(key='fit_mode__zh-CN').set_value('完整留白')
    app.selectbox(key='generic_palette__zh-CN').set_value('通用 32 色')
    app.slider(key='contrast').set_value(1.2).run()
    app.button(key='generate').click().run()
    assert not app.exception and not app.error
    original = app.session_state['result']
    old_indices = original['indices'].copy()
    old_settings_id = original['settings_id']
    for language in ('en', 'ja', 'zh-TW', 'zh-CN'):
        with (patch('core.prepare_image', side_effect=AssertionError('Image reprocessed')),
              patch('core.map_to_palette', side_effect=AssertionError('Image remapped'))):
            app.selectbox(key='language_choice').set_value(language).run()
        assert not app.exception and not app.error
        result = app.session_state['result']
        viewer = next(node for node in app.get('bidi_component') if node.proto.component_name == 'bead_grid_viewer')
        viewer_data = json.loads(viewer.proto.json)
        assert base64.b64decode(viewer_data['image'].split(',', 1)[1]) == result['grid']
        assert viewer_data['labels']['open'] == Translator(language)('放大查看')
        assert viewer_data['labels']['close'] == Translator(language)('关闭')
        assert app.query_params.get('lang') in (language, [language])
        assert app.title[0].value == Translator(language)('拼豆图生成器')
        assert result['settings_id'] == old_settings_id
        assert result['grid'] == original['grid'] and result['preview'] == original['preview']
        assert np.array_equal(result['indices'], old_indices)
        assert sum(item['count'] for item in result['usage']) == 13 * 9
        assert app.number_input(key='columns').value == 13
        assert app.number_input(key='rows').value == 9
        assert app.selectbox(key=f'fit_mode__{language}').value == '完整留白'
        assert app.selectbox(key=f'generic_palette__{language}').value == '通用 32 色'
        assert app.slider(key='contrast').value == 1.2
        assert len(app.get('download_button')) == 5
        with zipfile.ZipFile(BytesIO(result['bundle'])) as archive:
            assert archive.read('bead_pattern.xlsx') == result['xlsx']
            assert json.loads(archive.read('settings.json'))['export_language'] == language
        book = load_workbook(BytesIO(result['xlsx']))
        grid = book[Translator(language)('拼豆图')]
        assert grid.max_row == 9 and grid.max_column == 13
        for y, row in enumerate(grid):
            for x, cell in enumerate(row):
                assert cell.data_type == 's'
                assert cell.value == result['source_palette'][old_indices[y, x]].code
        pdf = PdfReader(BytesIO(result['pdf']))
        assert Translator(language)('拼豆图') in pdf.pages[0].extract_text()

    # A translated multi-select must retain a restricted real-brand palette.
    app.radio(key='palette_mode__zh-CN').set_value('内置品牌色卡').run()
    palette_id = app.selectbox(key='brand_palette_id__zh-CN').value
    groups_key = f'brand_groups_{palette_id}'
    selected_groups = app.multiselect(key=f'{groups_key}__zh-CN').value[:1]
    app.multiselect(key=f'{groups_key}__zh-CN').set_value(selected_groups).run()
    app.button(key='generate').click().run()
    assert not app.exception and not app.error
    original_brand = app.session_state['result']
    for language in ('ja', 'zh-TW', 'en', 'zh-CN'):
        app.selectbox(key='language_choice').set_value(language).run()
        assert not app.exception and not app.error
        result = app.session_state['result']
        assert app.multiselect(key=f'{groups_key}__{language}').value == selected_groups
        assert app.radio(key=f'palette_mode__{language}').value == '内置品牌色卡'
        assert app.selectbox(key=f'brand_palette_id__{language}').value == palette_id
        assert result['settings_id'] == original_brand['settings_id']
        assert result['source_palette'] == original_brand['source_palette']
        assert np.array_equal(result['indices'], original_brand['indices'])
        assert result['preview'] == original_brand['preview']


if __name__ == '__main__':
    unittest.main(verbosity=2)
