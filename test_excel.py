"""Run with: python -m unittest test_excel -v"""
import io
import unittest
import zipfile

import numpy as np
from openpyxl import load_workbook

from core import BeadColor, make_bundle
from excel_export import make_xlsx


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.palette = [BeadColor("001", "红色", 255, 0, 0, "品牌甲"),
                        BeadColor("B02", "白色", 255, 255, 255, "品牌乙"),
                        BeadColor("K03", "黑色", 0, 0, 0)]
        self.indices = np.array([[0, 1, 2], [2, 0, 1]], dtype=np.int16)

    def test_exact_pixel_mapping_and_statistics(self):
        book = load_workbook(io.BytesIO(make_xlsx(self.indices, self.palette)))
        self.assertEqual(book.sheetnames, ["拼豆图", "色号统计"])
        sheet = book["拼豆图"]
        self.assertEqual((sheet.max_row, sheet.max_column), (2, 3))
        self.assertFalse(sheet.merged_cells.ranges)
        for row in range(2):
            for col in range(3):
                cell = sheet.cell(row + 1, col + 1)
                bead = self.palette[int(self.indices[row, col])]
                self.assertEqual(cell.value, bead.code)
                self.assertEqual(cell.data_type, "s")
                self.assertEqual(cell.fill.fgColor.rgb[-6:], bead.hex[1:])
                self.assertEqual(cell.alignment.horizontal, "center")
        self.assertEqual(sheet["C1"].font.color.rgb[-6:], "FFFFFF")
        self.assertEqual(sheet["B1"].font.color.rgb[-6:], "16181C")
        self.assertAlmostEqual(sheet.column_dimensions["A"].width * 7 + 5,
                               sheet.row_dimensions[1].height / .75)
        stats = book["色号统计"]
        self.assertEqual(sum(stats.cell(row, 6).value for row in range(2, 5)), 6)
        self.assertEqual(stats["F5"].value, "=SUM(F2:F4)")
        self.assertEqual(stats["B2"].value, "品牌甲")

    def test_hidden_codes_still_have_every_colored_cell(self):
        book = load_workbook(io.BytesIO(make_xlsx(self.indices, self.palette, False)))
        sheet = book["拼豆图"]
        self.assertEqual((sheet.max_row, sheet.max_column), (2, 3))
        self.assertTrue(all(c.value is None for row in sheet for c in row))
        self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "FF0000")

    def test_imported_labels_are_literal_not_formulas(self):
        palette = [BeadColor("=1+1", "@name", 10, 20, 30, "=HYPERLINK(\"https://example.com\")")]
        book = load_workbook(io.BytesIO(make_xlsx(np.zeros((1, 1), dtype=int), palette)))
        for cell in (book["拼豆图"]["A1"], book["色号统计"]["B2"], book["色号统计"]["C2"]):
            self.assertEqual(cell.data_type, "s")
        self.assertEqual(book["拼豆图"]["A1"].value, "=1+1")

    def test_duplicate_codes_keep_distinct_palette_entries(self):
        palette = [self.palette[0], BeadColor("001", "蓝色", 0, 0, 255, "品牌乙")]
        book = load_workbook(io.BytesIO(make_xlsx(np.array([[0, 1, 1]]), palette)))
        stats = book["色号统计"]
        self.assertEqual({(stats.cell(r, 2).value, stats.cell(r, 6).value) for r in (2, 3)},
                         {("品牌甲", 1), ("品牌乙", 2)})

    def test_maximum_size_and_five_cell_borders(self):
        book = load_workbook(io.BytesIO(make_xlsx(np.zeros((300, 300), dtype=np.int16), self.palette)))
        sheet = book["拼豆图"]
        self.assertEqual((sheet.max_row, sheet.max_column), (300, 300))
        self.assertEqual(sheet["KN300"].value, "001")
        self.assertEqual(sheet["E1"].border.right.style, "medium")
        self.assertEqual(sheet["F1"].border.left.style, "medium")
        self.assertEqual(sheet["B2"].border.right.style, "thin")

    def test_invalid_input(self):
        for indices in (np.array([]), np.empty((0, 3), dtype=int), np.array([[-1]]),
                        np.array([[3]]), np.array([[.5]]), np.zeros((301, 1), dtype=int)):
            with self.subTest(shape=indices.shape), self.assertRaises(ValueError):
                make_xlsx(indices, self.palette)
        with self.assertRaises(ValueError):
            make_xlsx(np.zeros((1, 1), dtype=int), [])

    def test_bundle_contains_same_excel_and_old_call_still_works(self):
        excel = make_xlsx(self.indices, self.palette)
        args = (b"png", b"preview", b"csv", b"pdf", self.palette, {})
        with zipfile.ZipFile(io.BytesIO(make_bundle(*args, xlsx_bytes=excel))) as archive:
            self.assertEqual(archive.read("bead_pattern.xlsx"), excel)
        with zipfile.ZipFile(io.BytesIO(make_bundle(*args))) as archive:
            self.assertNotIn("bead_pattern.xlsx", archive.namelist())


if __name__ == "__main__":
    unittest.main()
