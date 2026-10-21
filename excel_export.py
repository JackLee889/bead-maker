"""Application-side Excel export; one converted bead is one editable cell."""
from __future__ import annotations

from i18n import LocalizedError, Translator

import io
from typing import TYPE_CHECKING, Sequence

import numpy as np
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from core import BeadColor


def _literal(cell, value: str) -> None:
    """Preserve leading zeros and never interpret imported labels as formulas."""
    if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
        raise LocalizedError('色号、品牌或名称含 Excel 不支持的控制字符，或长度超过 32767。')
    cell.value = value
    cell.data_type = "s"


def make_xlsx(
    indices: np.ndarray,
    palette: Sequence[BeadColor],
    show_codes: bool = True,
    *,
    language: str = "zh-CN",
) -> bytes:
    """Return XLSX bytes from the unscaled (rows, columns) palette-index array.

    Grid A1 is source pixel [0, 0]. No headers, merged cells or images occupy
    the grid. The legend is a generation-time snapshot, not a live inventory.
    """
    from core import _text_color, usage_rows
    t = Translator(language)

    indices = np.asarray(indices)
    if indices.ndim != 2 or not indices.size:
        raise LocalizedError('Excel 导出需要非空二维像素矩阵。')
    rows, columns = indices.shape
    if rows > 300 or columns > 300:
        raise LocalizedError('Excel 导出最多支持 300 × 300 格。')
    if not np.issubdtype(indices.dtype, np.integer):
        raise LocalizedError('像素矩阵必须包含整数色库索引。')
    if not palette or indices.min() < 0 or indices.max() >= len(palette):
        raise LocalizedError('像素矩阵含无效色库索引。')
    for color in palette:
        for label in (color.code, color.brand, color.name):
            if len(label) > 32767 or ILLEGAL_CHARACTERS_RE.search(label):
                raise LocalizedError('色号、品牌或名称含 Excel 不支持的控制字符，或长度超过 32767。')

    book = Workbook()
    grid = book.active
    grid.title = t("拼豆图")
    grid.sheet_view.showGridLines = False
    grid.sheet_view.zoomScale = 70
    # Calibri 11: width ~= (pixels - 5) / 7; points = pixels * 0.75.
    # Excel/WPS font and display scaling can introduce a small visual difference.
    for column in range(1, columns + 1):
        grid.column_dimensions[get_column_letter(column)].width = 27 / 7
    thin = Side(style="thin", color="9CA3AF")
    major = Side(style="medium", color="374151")
    borders = {}
    fills = [PatternFill("solid", fgColor=color.hex[1:]) for color in palette]
    fonts = [Font(name="Calibri", size=8, color="".join(f"{v:02X}" for v in _text_color(c.rgb))) for c in palette]
    alignment = Alignment(horizontal="center", vertical="center", shrinkToFit=True)
    for row in range(rows):
        grid.row_dimensions[row + 1].height = 24
        for column in range(columns):
            index = int(indices[row, column])
            cell = grid.cell(row + 1, column + 1)
            if show_codes:
                _literal(cell, palette[index].code)
            cell.fill = fills[index]
            cell.font = fonts[index]
            cell.alignment = alignment
            edges = (column % 5 == 0, row % 5 == 0,
                     (column + 1) % 5 == 0 or column == columns - 1,
                     (row + 1) % 5 == 0 or row == rows - 1)
            if edges not in borders:
                left, top, right, bottom = [major if edge else thin for edge in edges]
                borders[edges] = Border(left=left, top=top, right=right, bottom=bottom)
            cell.border = borders[edges]

    grid.print_area = f"A1:{get_column_letter(columns)}{rows}"
    grid.page_setup.paperSize = grid.PAPERSIZE_A4
    grid.page_setup.orientation = "landscape" if columns > rows else "portrait"
    grid.page_setup.scale = 100
    grid.print_options.headings = True

    legend = book.create_sheet(t("色号统计"))
    legend.append([t(label) for label in ("色块", "品牌", "色号", "名称", "HEX", "数量（生成时）")])
    usage = usage_rows(indices, palette)
    for row, item in enumerate(usage, 2):
        legend.cell(row, 1).fill = PatternFill("solid", fgColor=str(item["hex"])[1:])
        for column, key in enumerate(("brand", "code", "name", "hex"), 2):
            _literal(legend.cell(row, column), str(item[key]))
        legend.cell(row, 6, int(item["count"])).number_format = "#,##0"
    total_row = len(usage) + 2
    legend.cell(total_row, 5, t("合计"))
    legend.cell(total_row, 6, f"=SUM(F2:F{total_row - 1})").number_format = "#,##0"
    for cell in legend[1]:
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.font = Font(color="FFFFFF", bold=True)
    for column, width in zip("ABCDEF", (8, 24, 20, 28, 14, 20)):
        legend.column_dimensions[column].width = width
    for row in legend:
        legend.row_dimensions[row[0].row].height = 26
        for cell in row:
            cell.alignment = Alignment(vertical="center", shrinkToFit=True)
    legend.freeze_panes = "A2"
    legend.auto_filter.ref = f"A1:F{total_row - 1}"
    legend.sheet_properties.pageSetUpPr.fitToPage = True
    legend.page_setup.fitToWidth = 1
    legend.page_setup.fitToHeight = 0
    legend.print_title_rows = "1:1"
    legend.print_area = f"A1:F{total_row}"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
